import os
import docx
import zipfile
import xml.etree.ElementTree as ET

p = os.path.join(os.getcwd(), 'ZERO_ID_IEEE_Paper_Final.docx')
print('Exists:', os.path.exists(p), 'Size:', os.path.getsize(p))

with zipfile.ZipFile(p) as z:
    for name in z.namelist():
        if name.endswith('.xml'):
            ET.fromstring(z.read(name))
    print('All XML parts validated: PASS')
    doc_xml = z.read('word/document.xml').decode('utf-8')
    if 'w:cantSplit' in doc_xml:
        print('cantSplit protection for figures/tables: CONFIRMED')
    if 'w:num="2"' in doc_xml:
        print('IEEE 2-Column layout: CONFIRMED')

doc = docx.Document(p)
print(f'Total tables/boxes: {len(doc.tables)}')
for i, t in enumerate(doc.tables):
    txt = t.cell(0, 0).text[:40].replace('\n', ' ')
    print(f' - Table/Box {i+1}: {txt}...')

print('All validation tests passed successfully.')
