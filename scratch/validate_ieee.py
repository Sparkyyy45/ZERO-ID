import os
import docx
import zipfile
import xml.etree.ElementTree as ET

p = os.path.join(os.getcwd(), 'ZERO_ID_IEEE_Research_Paper_Submission.docx')
print('Exists:', os.path.exists(p), 'Size:', os.path.getsize(p))

with zipfile.ZipFile(p) as z:
    for name in z.namelist():
        if name.endswith('.xml'):
            ET.fromstring(z.read(name))
    print('All XML parts validated: PASS')
    doc_xml = z.read('word/document.xml').decode('utf-8')
    if 'w:num="2"' in doc_xml:
        print('IEEE 2-Column layout: ACTIVE')
    else:
        print('IEEE 2-Column layout: NOT FOUND')

doc = docx.Document(p)
text = ' '.join([p_el.text for p_el in doc.paragraphs])
words = len(text.split())
print(f'Word count: {words}')
print(f'Paragraphs: {len(doc.paragraphs)}')
print(f'Tables: {len(doc.tables)}')

if 'privakyc' in text.lower():
    print('PrivaKYC found: FAIL')
else:
    print('No PrivaKYC found: PASS')

# Check citations and section structure
sections = ['I. INTRODUCTION', 'II. RELATED WORK', 'III. SYSTEM ARCHITECTURE', 'IV. CIRCUIT DESIGN', 'V. SECURITY ANALYSIS', 'VI. EXPERIMENTAL EVALUATION', 'VII. REGULATORY COMPLIANCE', 'VIII. DISCUSSION', 'IX. CONCLUSION', 'REFERENCES']
for s in sections:
    found = any(s in p.text.upper() for p in doc.paragraphs)
    print(f'Section {s}: {"FOUND" if found else "MISSING"}')

# Verify references count
ref_count = sum(1 for p in doc.paragraphs if p.text.strip().startswith('[') and ']' in p.text)
print(f'Numbered IEEE References: {ref_count}')

print('Validation completed successfully.')
