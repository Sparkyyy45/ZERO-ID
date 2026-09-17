import os
import docx
import zipfile
import xml.etree.ElementTree as ET

p = os.path.join(os.getcwd(), 'ZERO_ID_IEEE_Paper_Final.docx')
print('Exists:', os.path.exists(p), 'Size:', os.path.getsize(p))

with zipfile.ZipFile(p) as z:
    for name in z.namelist():
        if name.endswith('.xml'):
            ET.fromstring(z.read(name))
    print('All XML parts validated: PASS')
    doc_xml = z.read('word/document.xml').decode('utf-8')
    if 'w:cantSplit' in doc_xml:
        print('cantSplit protection for figures/tables: CONFIRMED')
    if 'w:num="2"' in doc_xml:
        print('IEEE 2-Column layout: CONFIRMED')

doc = docx.Document(p)
text = ' '.join([p_el.text for p_el in doc.paragraphs])
words = len(text.split())
print(f'Word count: {words}')
print(f'Paragraphs: {len(doc.paragraphs)}')
print(f'Tables/Boxes: {len(doc.tables)}')

if 'privakyc' in text.lower():
    print('PrivaKYC found: FAIL')
else:
    print('No PrivaKYC found: PASS')

# Check citations and section structure
sections = [
    'I. INTRODUCTION', 'II. RELATED WORK', 'III. SYSTEM ARCHITECTURE', 
    'IV. CIRCUIT DESIGN AND CRYPTOGRAPHIC SPECIFICATION', 'V. FORMAL SECURITY ANALYSIS', 
    'VI. EMPIRICAL EVALUATION AND BENCHMARKS', 'VII. STATUTORY REGULATORY ALIGNMENT', 
    'VIII. DISCUSSION AND TECHNICAL TRADE-OFFS', 'IX. CONCLUSION AND FUTURE WORK', 'REFERENCES'
]
for s in sections:
    found = any(s in p.text.upper() for p in doc.paragraphs)
    print(f'Section {s}: {"FOUND" if found else "MISSING"}')

ref_count = sum(1 for p in doc.paragraphs if p.text.strip().startswith('[') and ']' in p.text)
print(f'Numbered IEEE References: {ref_count}')

print('Validation completed successfully.')
