import os
import zipfile
import docx
import xml.etree.ElementTree as ET

dir_path = os.path.join(os.getcwd(), 'final docs')
files = [
    "ZERO_ID_IEEE_Paper_Final_Revised.docx",
    "ZERO_ID_Technical_Report_Patent_Specification_REVISED.docx"
]

report = {}

for fname in files:
    fpath = os.path.join(dir_path, fname)
    info = {"name": fname, "exists": os.path.exists(fpath)}
    if not info["exists"]:
        report[fname] = info
        continue
    
    info["size_bytes"] = os.path.getsize(fpath)
    
    # 1. XML Integrity check
    try:
        with zipfile.ZipFile(fpath, 'r') as z:
            for n in z.namelist():
                if n.endswith('.xml'):
                    ET.fromstring(z.read(n))
            doc_xml = z.read('word/document.xml').decode('utf-8')
            info["xml_valid"] = True
            info["has_two_col"] = 'w:num="2"' in doc_xml
            info["has_cantsplit"] = 'w:cantSplit' in doc_xml
    except Exception as e:
        info["xml_valid"] = False
        info["xml_error"] = str(e)
    
    # 2. Document Content Extraction
    try:
        doc = docx.Document(fpath)
        paragraphs = [p.text for p in doc.paragraphs]
        full_text = " ".join(paragraphs)
        info["word_count"] = len(full_text.split())
        info["paragraph_count"] = len(paragraphs)
        info["table_count"] = len(doc.tables)
        
        # Check for legacy branding
        info["has_priva"] = "privakyc" in full_text.lower()
        
        # Check headings
        headings = [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith(('I.', 'II.', 'III.', 'IV.', 'V.', 'VI.', 'VII.', 'VIII.', 'IX.', '1.', '2.', '3.', '4.', '5.', 'CLAIM', 'Abstract', 'REFERENCES'))]
        info["headings_sample"] = headings[:12]
        
        # Check math notation quality
        math_check = {
            "has_raw_latex": any(t in full_text for t in ['Adv_{', 'Exp_{', '_{A}', '2^{-lambda}', '\\lambda', 'approx_c']),
            "has_unicode_math": any(t in full_text for t in ['Adv_A', 'bn254_pairing', 'e(A, B)', 'NullifierHash', 'Poseidon', 'Groth16', '≤', '≈', '∈', 'λ', 'π', 'σ'])
        }
        info["math_check"] = math_check
        
        # Check tables info
        table_summaries = []
        for i, t in enumerate(doc.tables):
            try:
                hdr = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells]
                table_summaries.append(f"Table {i+1}: {len(t.rows)} rows x {len(t.columns)} cols -> Header: {hdr}")
            except Exception as te:
                table_summaries.append(f"Table {i+1}: Error reading ({te})")
        info["tables_info"] = table_summaries
        
    except Exception as e:
        info["doc_read_error"] = str(e)
        
    report[fname] = info

print("=== VERIFICATION RESULTS ===")
for fname, res in report.items():
    print(f"\n--- {fname} ---")
    for k, v in res.items():
        if k != "headings_sample" and k != "tables_info":
            print(f"  {k}: {v}")
    print("  Headings:")
    for h in res.get("headings_sample", []):
        print(f"    - {h}")
    print("  Tables:")
    for t in res.get("tables_info", []):
        print(f"    - {t}")
